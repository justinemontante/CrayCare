"""Finish WQAD dictionary alignment (remaining deltas after partial update).
Table 16: delete model_version (never written); clarify detection_id as doc
  ID; fix source + driver formats; normalize timestamp row.
Table 17: clarify contributors[] is an embedded array, not a collection."""

from docx import Document

SRC = "docs/Deliverables/Data Dictionary.docx"


def set_cell(cell, text):
    if cell.text.strip() == text.strip():
        return
    cell.text = ""
    cell.paragraphs[0].add_run(text)


doc = Document(SRC)
assert len(doc.tables) == 21, len(doc.tables)
t16, t17 = doc.tables[15], doc.tables[16]

names16 = [r.cells[0].text.strip() for r in t16.rows[1:]]
assert "model_version" in names16, names16

for row in list(t16.rows[1:]):
    name = row.cells[0].text.strip()
    if name == "model_version":
        row._tr.getparent().remove(row._tr)
    elif name == "detection_id":
        set_cell(
            row.cells[4],
            'Document ID of the detection record ("current" live alias or '
            "YYYYMMDDTHHMMSS history ID); not a stored field.",
        )
    elif name == "source":
        set_cell(
            row.cells[3],
            "WQAD model | Model unavailable | Insufficient data | Stale sensor history",
        )
        set_cell(
            row.cells[4],
            "Detection origin: model result, or the reason a model result "
            "was unavailable.",
        )
    elif name == "driver":
        set_cell(
            row.cells[3],
            "temp | pH | DO | turbidity | waterLevel | overall | N/A",
        )
    elif name == "timestamp":
        set_cell(row.cells[1], "timestamp")
        set_cell(row.cells[3], "2026-09-06T08:00:00+00:00")
        set_cell(
            row.cells[4],
            "Processing time in ISO-8601; read by the app (with processed_at) "
            "for display and ordering.",
        )

for row in t17.rows[1:]:
    name = row.cells[0].text.strip()
    if name == "detection_id":
        set_cell(
            row.cells[4],
            "ID of the parent detection document; contributors[] is embedded "
            "in that document, not a separate collection.",
        )
    elif name == "rank_no":
        set_cell(
            row.cells[4],
            "Position in the embedded contributors[] array (max 3 elements); "
            "not a separate table in Firestore.",
        )

doc.save(SRC)
print("saved", SRC)
