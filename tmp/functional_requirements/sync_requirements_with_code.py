from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
DOC_PATH = ROOT / "docs" / "Deliverables" / "Functional and Non-Functional Requirements.docx"


def set_cell(cell, value):
    paragraph = cell.paragraphs[0]
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = value
    else:
        paragraph.add_run(value)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def replace_row(table, starts_with, values):
    for row in table.rows:
        if row.cells[0].text.strip().startswith(starts_with):
            for cell, value in zip(row.cells, values):
                set_cell(cell, value)
            return
    raise RuntimeError(f"Missing requirement row: {starts_with}")


doc = Document(DOC_PATH)

# Hardware/functional matrix: include the newly integrated feed-level sensor.
hardware_continuation = doc.tables[1]
if not any(
    row.cells[0].text.strip().startswith("The system must monitor the feeder hopper")
    for row in hardware_continuation.rows
):
    template_row = hardware_continuation.rows[3]._tr
    new_row = deepcopy(template_row)
    hardware_continuation.rows[4]._tr.addprevious(new_row)
    added = hardware_continuation.rows[4]
    for cell, value in zip(
        added.cells,
        (
            "The system must monitor the feeder hopper level in real time",
            "Feed-level sensing through the ESP32, including percentage and estimated remaining grams based on the configured hopper capacity",
            "High",
            "Implemented",
        ),
    ):
        set_cell(cell, value)

replace_row(
    hardware_continuation,
    "The system must support role-based user access",
    (
        "The system must support role-based user access",
        "Role-based access control for the single System Admin and registered Tank Owners, with permissions enforced in the application and Firebase Security Rules",
        "High",
        "Implemented",
    ),
)
replace_row(
    hardware_continuation,
    "The system must restrict write actions",
    (
        "The system must restrict tank-operation write actions to the designated owner",
        "Owner-only threshold, feeding, actuator, sampling, and tank-operation writes; the System Admin manages accounts and hardware assignment",
        "High",
        "Implemented",
    ),
)
replace_row(
    hardware_continuation,
    "The system must manage user accounts",
    (
        "The system must allow the administrator to manage owner accounts and hardware assignment",
        "Administrative interface for viewing owner accounts, enabling or disabling accounts, and safely assigning or reassigning the ESP32 hardware to one initialized owner tank",
        "High",
        "Implemented",
    ),
)

software_matrix = doc.tables[2]
replace_row(
    hardware_continuation,
    "The system must record initial stock setup",
    (
        "The system must record the initial tank and baseline sampling setup",
        "Tank initialization form for stocking date, initial population, fixed batch sample size, total baseline sample weight, and total baseline sample length; the baseline sampling record is saved with the new active batch",
        "High",
        "Implemented",
    ),
)
replace_row(
    software_matrix,
    "The system must classify growth stages",
    (
        "The system must present the current growth stage",
        "Automatic presentation of Early Juvenile, Advanced Juvenile, Pre-Adult, or Market Size using the latest valid Average Body Weight or Average Body Length measurement; sensor thresholds remain separately configured by the owner",
        "Medium",
        "Implemented",
    ),
)
replace_row(
    software_matrix,
    "The system must analyze historical and real-time parameters",
    (
        "The system must analyze recent multivariate water-sensor patterns using machine learning to detect unusual conditions",
        "Scheduled Water Quality Anomaly Detection using an unsupervised Isolation Forest model over temperature, pH, dissolved oxygen, turbidity, water level, and recent trend features; results are Normal, Unusual, or Insufficient",
        "High",
        "Implemented",
    ),
)

software_continuation = doc.tables[3]
if not any(
    row.cells[0].text.strip().startswith("The system must generate and download reports")
    for row in software_continuation.rows
):
    report_template = software_continuation.rows[2]._tr
    report_row_xml = deepcopy(report_template)
    software_continuation.rows[3]._tr.addprevious(report_row_xml)
    report_row = software_continuation.rows[3]
    for cell, value in zip(
        report_row.cells,
        (
            "The system must generate and download reports",
            "PDF and CSV export for growth and production records and Water Quality Anomaly Detection history, using the selected reporting period and the same validated calculations shown in the application",
            "Medium",
            "Implemented",
        ),
    ):
        set_cell(cell, value)

replace_row(
    software_continuation,
    "The system must cache sensor data locally",
    (
        "The system must preserve sensor information during internet interruptions",
        "The mobile application displays the last cached readings through local and Firebase offline persistence, while the ESP32 buffers sensor history and synchronizes it after connectivity returns",
        "Medium",
        "Implemented",
    ),
)

non_functional_continuation = doc.tables[5]
if not any(
    row.cells[0].text.strip().startswith("The system should generate consistent and readable reports")
    for row in non_functional_continuation.rows
):
    nfr_template = non_functional_continuation.rows[-1]._tr
    nfr_row_xml = deepcopy(nfr_template)
    non_functional_continuation._tbl.append(nfr_row_xml)
    nfr_row = non_functional_continuation.rows[-1]
    for cell, value in zip(
        nfr_row.cells,
        (
            "The system should generate consistent and readable reports",
            "PDF and CSV exports preserve the selected date range, field labels, units, record values, and calculations consistently without modifying the stored source data",
            "Medium",
            "Implemented",
        ),
    ):
        set_cell(cell, value)

hardware_non_functional = doc.tables[4]
for requirement in (
    "The system should operate continuously in real-time farm conditions",
    "The system should ensure reliable sensor readings",
    "The system should withstand aquaculture environment conditions",
    "The system should support continuous data transmission",
):
    for row in hardware_non_functional.rows:
        if row.cells[0].text.strip().startswith(requirement):
            set_cell(row.cells[3], "Implemented - For Field Validation")
            break
    else:
        raise RuntimeError(f"Missing hardware non-functional requirement: {requirement}")
replace_row(
    software_continuation,
    "The system must generate machine learning-based insights",
    (
        "The system must explain Water Quality Anomaly Detection results and provide verification-focused recommendations",
        "Isolation Forest WQAD presents an anomaly score, primary contributing sensor, ranked contributors, cautious insight, and recommended checks without claiming a confirmed cause or controlling devices",
        "High",
        "Implemented",
    ),
)

doc.save(DOC_PATH)

# Content assertions against the current implementation contract.
verified = Document(DOC_PATH)
text = "\n".join(
    cell.text
    for table in verified.tables
    for row in table.rows
    for cell in row.cells
)
required = (
    "Isolation Forest",
    "Normal, Unusual, or Insufficient",
    "Feed-level sensing",
    "single System Admin",
    "ESP32 buffers sensor history",
    "generate and download reports",
    "consistent and readable reports",
    "baseline sampling setup",
    "sensor thresholds remain separately configured",
)
outdated = ("Random Forest", "Hugging Face", "Monitor User", "predict water quality health")
missing = [item for item in required if item not in text]
stale = [item for item in outdated if item in text]
if missing or stale:
    raise RuntimeError(f"Sync verification failed: missing={missing}; stale={stale}")

print(f"Updated: {DOC_PATH}")
print(f"Tables: {len(verified.tables)}")
print("Current implementation terminology verified.")
