from copy import deepcopy
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "Deliverables" / "Chapters 1–3.docx"
OUTPUT = ROOT / "docs" / "Deliverables" / "Use Case Report.docx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def text_of(element):
    return " ".join(element.xpath(".//w:t/text()", namespaces=NS)).strip()


def set_cell(cell, value):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def find_use_case(doc, name):
    for table in doc.tables:
        for row in table.rows[1:]:
            if row.cells[0].text.strip() == name:
                return table, row
    raise RuntimeError(f"Use case not found: {name}")


def update_use_case(doc, name, values):
    _table, row = find_use_case(doc, name)
    for cell, value in zip(row.cells, values):
        set_cell(cell, value)


# Preserve the complete source package, page setup, header, footer, styles, and
# relationships; only retain the Use Case Report portion of Chapter 3.
with ZipFile(SOURCE, "r") as source_zip:
    root = etree.fromstring(source_zip.read("word/document.xml"))
    body = root.find("w:body", NS)
    children = list(body)
    start = next(
        i for i, child in enumerate(children)
        if text_of(child).strip().upper() == "3.1.3 USE CASE REPORT"
    )
    end = next(
        i for i, child in enumerate(children[start + 1 :], start + 1)
        if text_of(child).strip().upper().startswith("3.2 DESIGN SPECIFICATIONS")
    )
    section = next(
        (deepcopy(child) for child in reversed(children) if child.tag == f"{{{W}}}sectPr"),
        None,
    )
    selected = [deepcopy(child) for child in children[start:end]]
    for child in list(body):
        body.remove(child)
    for child in selected:
        body.append(child)
    if section is not None:
        body.append(section)
    updated_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(OUTPUT, "w", ZIP_DEFLATED) as output_zip:
        for item in source_zip.infolist():
            payload = updated_xml if item.filename == "word/document.xml" else source_zip.read(item.filename)
            output_zip.writestr(item, payload)


doc = Document(OUTPUT)

# Authentication is a current user-facing workflow and must be represented in
# the report rather than being implied only by other use-case preconditions.
auth_table, auth_anchor = find_use_case(doc, "View Real-Time Dashboard")
auth_row_xml = deepcopy(auth_anchor._tr)
auth_anchor._tr.addprevious(auth_row_xml)
auth_row = auth_table.rows[1]
for cell, value in zip(auth_row.cells, (
    "Register and Authenticate User Account",
    "Farm Owner, System Admin",
    "Allows a new owner to register using email and password or Google, verify an email-based account, sign in securely, request a password-reset link, and sign out. The existing System Admin can sign in and sign out but cannot create another admin account.",
    "The mobile application must be installed and have internet access for online authentication. Owner registration requires valid account information. Admin access requires the existing active admin account.",
    "1. A new owner selects email registration or Google sign-in; an email/password owner completes email verification when required.\n2. A registered owner or the System Admin signs in using the supported credential method.\n3. The authentication service validates the credential and loads the matching user profile and role.\n4. If the user forgets an email-account password, the user requests a reset link from the login screen.\n5. An authenticated user may sign out from Profile and Settings.",
    "A valid active account is routed to the correct owner or admin interface. Registration creates an owner profile only. Signing out ends the current application session.",
    "Invalid credentials, unverified email, disabled account, unavailable internet, cancelled Google selection, or authentication-service failure prevents access and displays an applicable message.",
)):
    set_cell(cell, value)

update_use_case(doc, "View Real-Time Dashboard", (
    "View Real-Time Dashboard",
    "Farm Owner",
    "Views live water-sensor readings, feed level, connectivity, tank production information, sampling countdown, and feeding status from the dashboard.",
    "The owner must be authenticated. The owner tank must be initialized for production information to appear.",
    "1. The owner opens the Dashboard.\n2. The system loads the assigned tank and latest routed ESP32 readings.\n3. The system displays temperature, pH, dissolved oxygen, turbidity, and water-level cards with status indicators.\n4. The system displays feed level, tank production information, next sampling, and feeding progress.\n5. The owner may open a sensor or feed-level detail panel or use a quick action.",
    "The latest available monitoring and operational information is displayed and updates when new data arrives.",
    "If the tank is not initialized, the setup state is shown. If the ESP32 is offline or data is stale, the interface shows the applicable connection or waiting message and retains the latest cached values when available.",
))

update_use_case(doc, "Monitor Water Parameters", (
    "Monitor Water Parameters",
    "Farm Owner",
    "Monitors temperature, pH, dissolved oxygen, turbidity, and water level through live cards, status labels, ideal ranges, and detailed information panels.",
    "The owner must be authenticated and have access to an initialized tank. Sensor readings and threshold settings must be available.",
    "1. The system receives the latest routed readings from the assigned ESP32.\n2. It validates and displays each water parameter, unit, status, and ideal range.\n3. The owner taps a parameter card.\n4. The system opens the corresponding detail panel with the current value and configured threshold interpretation.",
    "The owner can identify the latest water conditions and review the applicable configured limits.",
    "Invalid or missing sensor data is shown as unavailable. Stale readings do not appear as current live measurements.",
))

update_use_case(doc, "Control Hardware Devices", (
    "Control Hardware Devices",
    "Farm Owner",
    "Controls Aerator 1, Aerator 2, and the water pump using OFF, AUTO, or ON mode and reviews device activity logs.",
    "The owner must be authenticated, the tank must be initialized, and the hardware must be actively assigned to that owner tank.",
    "1. The owner opens Controls and selects Actuators.\n2. The system displays each device and its current mode.\n3. The owner selects OFF, AUTO, or ON.\n4. The command is saved for the assigned ESP32.\n5. The resulting device activity is recorded and displayed in the log history.",
    "The selected control mode is stored and the assigned ESP32 can apply it; the activity log reflects the operation.",
    "If the hardware is unassigned, assigned to another tank, offline, or the write fails, the command is blocked or reported as unsuccessful.",
))

update_use_case(doc, "Manage Feeding System", (
    "Manage Feeding System",
    "Farm Owner",
    "Creates recurring feeding schedules, triggers Feed Now, checks hopper sufficiency, and reviews ESP-confirmed feeding outcomes and daily consumption.",
    "The owner must be authenticated, the tank initialized, and the feeder assigned to the owner tank. Feed-level and hopper-capacity information should be available for sufficiency checks.",
    "1. The owner opens the Feeding tab.\n2. The owner adds or edits a time, repeat days, and feed amount; the system rejects overlapping day-and-time combinations.\n3. For Feed Now, the system warns about nearby schedules and blocks a conflicting immediate execution window.\n4. Before dispensing, the ESP32 checks the estimated available feed.\n5. The feeder records checking, dispensing, completed, skipped-insufficient, blocked, or failed status with relevant feed-level details.",
    "Valid schedules are saved. Manual and scheduled outcomes update feeding history, next feeding, progress, and completed-feed consumption totals.",
    "Duplicate schedules are rejected. Feeding is skipped when the hopper is empty, feed data is unavailable, or estimated feed is insufficient. Offline or failed dispensing is recorded with an applicable reason.",
))

update_use_case(doc, "Receive Alerts & Notifications", (
    "Receive Alerts and Notifications",
    "Farm Owner",
    "Receives and reviews water-condition alerts, device events, feeding results, reminders, and feed-level notifications.",
    "The owner must be authenticated. System notification permission and the applicable preference must be enabled for push delivery.",
    "1. A supported sensor, feeder, actuator, or reminder event occurs.\n2. The backend verifies the active owner-tank relationship and notification preference.\n3. The system stores the notification and sends an FCM push message when permitted.\n4. The owner opens Notifications and filters or reviews the grouped records.",
    "The notification is available in the owner account, and push delivery is attempted when device permission and preferences allow it.",
    "When the device is offline or push permission is disabled, delivery may be delayed or limited to the in-app notification history.",
))

update_use_case(doc, "View Analytics & Historical Data", (
    "View Analytics and Historical Data",
    "Farm Owner",
    "Views historical water-sensor charts and minimum, maximum, and average values across Live, 24H, 7D, 30D, and Custom ranges.",
    "The owner must be authenticated and historical sensor readings must exist for the selected tank and period.",
    "1. The owner opens Analytics.\n2. The owner selects a time range.\n3. The system queries the relevant history records.\n4. It displays the available data points, charts, and summary statistics.\n5. The owner may open a larger chart view.",
    "The selected historical range and its available statistics are displayed.",
    "If no valid data exists for the range, the system displays an empty-data state instead of fabricating points.",
))

update_use_case(doc, "Manage Tank Grow-out Records", (
    "Manage Tank Grow-out Records",
    "Farm Owner",
    "Initializes a production batch and records fixed-size weekly sampling, mortality, and harvest information for the current batch.",
    "The owner must be authenticated. Initialization is required before later production records. After the baseline sampling, its sample size is retained for succeeding samples in the same batch.",
    "1. The owner initializes the tank with stock count and stocking date.\n2. The owner records the baseline sample count, total weight, and total length.\n3. On an eligible later calendar date, the owner records the next sample using the fixed batch sample size.\n4. The system computes ABW, ABL, estimated biomass, survival, and growth summaries.\n5. Mortality and harvest records update the live population without exceeding the available count.",
    "Validated production records are saved for the current batch and the overview and trend displays are updated.",
    "Invalid dates, duplicate baseline initialization, inconsistent sample size after baseline, or mortality/harvest counts exceeding the live population are rejected.",
))

update_use_case(doc, "View Growth Trends & Tank Records", (
    "View Growth Trends and Tank Records",
    "Farm Owner",
    "Reviews sampling-based ABW and ABL trends, estimated biomass, mortality, survival, current growth stage, and batch records.",
    "The owner must be authenticated and the tank must have an initialized batch. Trend comparisons require sufficient sampling records.",
    "1. The owner opens Tank Production.\n2. The system loads the current batch inventory, sampling, mortality, and harvest records.\n3. It calculates and displays sampling-based growth measures, estimated biomass, survival, and the configured growth-stage presentation.\n4. The owner switches among the tank tabs to review details and history.",
    "The current batch overview and available production trends are displayed from saved records.",
    "When no batch or insufficient sampling history exists, the applicable setup or insufficient-data state is shown.",
))

update_use_case(doc, "Manage Water Thresholds", (
    "Manage Sensor Thresholds",
    "Farm Owner",
    "Updates water-sensor limits and feed-level settings used by dashboard status, alerts, and rule-based safety behavior.",
    "The owner must be authenticated and have an initialized tank.",
    "1. The owner opens Sensor Thresholds in Settings.\n2. The owner edits valid minimum and maximum values for a water sensor or the low, critical, and hopper-capacity settings for feed level.\n3. The system validates the values.\n4. The owner confirms Update.\n5. The settings are stored under the owner tank.",
    "Valid threshold settings are saved and reflected by dependent dashboard, notification, and control logic.",
    "Invalid, reversed, non-finite, or out-of-range values are rejected and the previous valid configuration is retained.",
))

update_use_case(doc, "Manage Personal Account", (
    "Manage Personal Account",
    "Farm Owner, System Admin",
    "Updates permitted profile information and manages password credentials according to the account sign-in method.",
    "The user must be authenticated and be viewing Profile and Settings.",
    "1. The user opens Edit Profile or Change Password.\n2. The user supplies valid information and any required reauthentication.\n3. The authentication service and profile record are updated.\n4. The interface confirms completion.",
    "The permitted account information or password credential is updated.",
    "Invalid input, a recent-login requirement, an authentication-provider limitation, or a network failure prevents the update and displays a clear message.",
))

update_use_case(doc, "Manage Notification Preferences", (
    "Manage Notification Preferences",
    "Farm Owner",
    "Configures supported CrayCare notification categories and sound or vibration behavior.",
    "The owner must be authenticated and viewing Notification Settings.",
    "1. The owner opens Notification Settings.\n2. The system loads the saved owner preferences.\n3. The owner changes a supported preference.\n4. The setting is saved and applied to future eligible notifications.",
    "Future notification processing follows the saved preferences, subject to operating-system permission.",
    "If device-level notification permission is disabled, the application directs the owner to the operating-system settings when appropriate.",
))

update_use_case(doc, "Manage User Accounts", (
    "Manage Owner Accounts and Hardware Assignment",
    "System Admin",
    "Views owner accounts, enables or disables them, and assigns or reassigns the single ESP32 hardware system to one initialized active owner tank.",
    "The administrator must be authenticated with the active admin account. The target owner must be active and must have completed tank initialization before assignment.",
    "1. The admin opens the account list and searches or filters owners.\n2. The admin may enable or disable an owner account.\n3. For assignment, the system validates the target owner and tank.\n4. If another owner is assigned, the system warns that reassignment changes the data route.\n5. After confirmation, the assignment metadata is updated atomically.",
    "Account status or the active owner-tank hardware assignment is updated. New ESP32 data is routed only to the valid assigned tank.",
    "The admin cannot disable the admin account, convert an owner into another admin, assign hardware to a disabled/uninitialized owner, or silently reassign hardware without confirmation.",
))

# Add the current machine-learning use case after Analytics, using the same row style.
analytics_table, analytics_row = find_use_case(doc, "View Analytics and Historical Data")
new_row = deepcopy(analytics_row._tr)
analytics_row._tr.addnext(new_row)
added_row = analytics_table.rows[2]
for cell, value in zip(added_row.cells, (
    "Review Water Quality Anomaly Detection",
    "Farm Owner",
    "Reviews the latest and historical WQAD results produced by the unsupervised Isolation Forest model, including status, anomaly score, contributing sensors, cautious insight, and recommended checks.",
    "The owner must be authenticated. The tank must have at least twelve contiguous valid ten-minute readings covering approximately two hours, and a deployed WQAD model must be available.",
    "1. The hourly backend validates the active hardware assignment and recent sensor window.\n2. It derives multivariate level, spread, change, and trend features without using safety thresholds as model inputs or labels.\n3. Isolation Forest returns a Normal or Unusual pattern result; incomplete or stale data produces Insufficient.\n4. The system stores the result.\n5. The owner opens the WQAD card or history and reviews contributors, insight, and recommended verification steps.",
    "A WQAD result and explanatory decision-support information are available to the owner. WQAD does not directly control devices or claim a confirmed biological cause.",
    "Missing, stale, discontinuous, or malformed readings return Insufficient. Prototype findings remain advisory until the model is retrained and validated using calibrated field data.",
)):
    set_cell(cell, value)

# Add report export as a distinct user-visible function after WQAD review.
report_anchor_table, report_anchor_row = find_use_case(
    doc, "Review Water Quality Anomaly Detection"
)
report_row_xml = deepcopy(report_anchor_row._tr)
report_anchor_row._tr.addnext(report_row_xml)
report_row = report_anchor_table.rows[3]
for cell, value in zip(report_row.cells, (
    "Generate and Download Reports",
    "Farm Owner",
    "Generates downloadable growth and production reports and Water Quality Anomaly Detection history in PDF or CSV format for review and research documentation.",
    "The owner must be authenticated. A selected production batch or available WQAD history must contain valid records for the requested report.",
    "1. The owner opens the export action from Growth Trends or WQAD History.\n2. The owner selects PDF or CSV.\n3. For growth reporting, the system uses the currently selected batch; for WQAD reporting, it uses the available anomaly-detection history.\n4. The system applies the same validated calculations and labels used by the application.\n5. The generated file is presented through the device share sheet for saving or sharing.",
    "A readable PDF or structured CSV report is generated from the current selected-batch or WQAD-history records without changing the source data.",
    "If there are no valid records, the system displays an empty-data message. File creation, storage, or sharing failures are reported without producing a misleading completed report.",
)):
    set_cell(cell, value)

# Update the closing explanation so it describes the two current roles.
for paragraph in doc.paragraphs:
    if "Table 3.3" in paragraph.text and "shows the Use Case Report" in paragraph.text:
        paragraph.text = (
            "Table 3.3 shows the Use Case Report of CrayCare. It presents the interactions of the "
            "Farm Owner and System Admin with the current monitoring, feeding, control, production, "
            "notification, analytics, Water Quality Anomaly Detection, account, and hardware-assignment features."
        )
    elif paragraph.text.startswith("The Use Case Report serves as"):
        paragraph.text = (
            "The Use Case Report explains the conditions, actions, expected results, and exceptions for "
            "each current system operation. It serves as a functional reference for checking whether the "
            "implemented application, Firebase backend, and assigned ESP32 behavior match the intended workflow."
        )

# The source table cells use Word automatic numbering. Since the synchronized
# content already contains deliberate inline numbering in Flow of Events,
# retaining w:numPr would display duplicate numbers (for example, "1. 1.").
# Remove inherited list numbering from every use-case table paragraph and drop
# empty continuation paragraphs left by the original multi-paragraph cells.
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in list(cell.paragraphs):
                p_pr = paragraph._p.pPr
                if p_pr is not None and p_pr.numPr is not None:
                    p_pr.remove(p_pr.numPr)
            for paragraph in list(cell.paragraphs[1:]):
                if not paragraph.text.strip():
                    cell._tc.remove(paragraph._p)

# Restore a single visible manual sequence in the list-oriented columns.
# Flow of Events already has its own 1..N sequence; the other three columns
# receive 1..N numbering based on their complete sentences.
for table in doc.tables:
    for row in table.rows[1:]:
        for column_index in (3, 5, 6):
            cell = row.cells[column_index]
            plain_text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            items = [item.strip() for item in re.split(r"(?<=[.!?])\s+", plain_text) if item.strip()]
            if items:
                set_cell(cell, "\n".join(f"{index}. {item}" for index, item in enumerate(items, 1)))

# Justified text creates irregular word spacing inside narrow report columns.
# Keep headers centered, but use flush-left body text and remove inherited list
# indents so every visible number begins at the same position.
for table in doc.tables:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if row_index == 0
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                paragraph.paragraph_format.left_indent = Inches(0)
                paragraph.paragraph_format.first_line_indent = Inches(0)

# Use one fixed grid for every continuation of Table 3.3. The total width is
# 8.30 inches, matching the usable width of the source landscape section.
column_widths = tuple(
    Inches(value)
    for value in (0.95, 0.70, 1.15, 1.25, 2.00, 1.15, 1.10)
)
for table in doc.tables:
    table.autofit = False
    grid_columns = table._tbl.tblGrid.gridCol_lst
    for column_index, width in enumerate(column_widths):
        table.columns[column_index].width = width
        if column_index < len(grid_columns):
            grid_columns[column_index].set(qn("w:w"), str(width.twips))
        for cell in table.columns[column_index].cells:
            cell.width = width

doc.save(OUTPUT)

# Structural and terminology verification.
verified = Document(OUTPUT)
all_text = "\n".join(
    [p.text for p in verified.paragraphs]
    + [cell.text for table in verified.tables for row in table.rows for cell in row.cells]
)
required = (
    "Review Water Quality Anomaly Detection",
    "Isolation Forest",
    "Feed-level",
    "System Admin",
    "Farm Owner",
    "hardware assignment",
    "Generate and Download Reports",
    "Register and Authenticate User Account",
)
outdated = ("Monitor User", "changes user role", "Random Forest", "Hugging Face")
missing = [value for value in required if value not in all_text]
stale = [value for value in outdated if value in all_text]
if missing or stale or len(verified.tables) != 9:
    raise RuntimeError(
        f"Verification failed: missing={missing}; stale={stale}; tables={len(verified.tables)}"
    )

numbered_paragraphs = [
    paragraph
    for table in verified.tables
    for row in table.rows
    for cell in row.cells
    for paragraph in cell.paragraphs
    if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
]
if numbered_paragraphs:
    raise RuntimeError(f"Inherited Word numbering remains in {len(numbered_paragraphs)} table paragraphs")

for table in verified.tables:
    for row in table.rows[1:]:
        for column_index in (3, 4, 5, 6):
            if not row.cells[column_index].text.strip().startswith("1. "):
                raise RuntimeError(
                    f"Visible numbering missing in column {column_index + 1}: {row.cells[0].text}"
                )

expected_widths = [width.twips for width in column_widths]
for table_number, table in enumerate(verified.tables, 1):
    actual_widths = [cell.width.twips for cell in table.rows[0].cells]
    if actual_widths != expected_widths:
        raise RuntimeError(
            f"Table {table_number} column widths differ: {actual_widths} != {expected_widths}"
        )

print(f"Created: {OUTPUT}")
print(f"Use case tables: {len(verified.tables)}")
print(f"Use case rows: {sum(max(0, len(t.rows) - 1) for t in verified.tables)}")
