from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

source = Path(r"docs/Deliverables/Data Dictionary - Updated.docx")
output = Path(r"docs/Deliverables/Data Dictionary - Normalized.docx")
doc = Document(source)

remove_from_tanks = {
    "stocking_date",
    "last_sample_date",
    "sample_count",
    "initial_total_sample_weight",
    "initial_total_sample_length",
    "initial_population",
}

# Tables follow their captions in the same order. Validate the expected third
# caption before changing the corresponding third table.
table_titles = [p.text.strip() for p in doc.paragraphs if p.text.strip().endswith(" Table")]
if len(table_titles) < 3 or table_titles[2] != "tanks Table":
    raise RuntimeError("Could not locate tanks Table")
tanks_table = doc.tables[2]

for row in list(tanks_table.rows)[1:]:
    if row.cells[0].text.strip() in remove_from_tanks:
        row._element.getparent().remove(row._element)

# Keep table captions sequential after the previously removed derived table.
counter = 0
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text.startswith("Table ") and text[6:].isdigit():
        counter += 1
        paragraph.text = f"Table {counter}"
        paragraph.alignment = 1
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

doc.save(output)
