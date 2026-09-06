import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfReader


docx_path = Path(sys.argv[1])
pdf_path = Path(sys.argv[2])
out_dir = Path(sys.argv[3])
out_dir.mkdir(parents=True, exist_ok=True)

reader = PdfReader(str(pdf_path))
pages = []
for index, page in enumerate(reader.pages, start=1):
    text = page.extract_text() or ""
    pages.append({"page": index, "text": text})

(out_dir / "pages.json").write_text(
    json.dumps(pages, ensure_ascii=False, indent=2), encoding="utf-8"
)
(out_dir / "full_text.txt").write_text(
    "\n\n".join(f"===== PAGE {p['page']} =====\n{p['text']}" for p in pages),
    encoding="utf-8",
)

doc = Document(str(docx_path))
paragraphs = []
style_counts = Counter()
for index, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    style = para.style.name if para.style else ""
    style_counts[style] += 1
    if text:
        paragraphs.append({"index": index, "style": style, "text": text})

tables = []
for table_index, table in enumerate(doc.tables):
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    tables.append({"index": table_index, "rows": rows})

sections = []
for index, section in enumerate(doc.sections):
    sections.append(
        {
            "index": index,
            "width_inches": round(section.page_width.inches, 3),
            "height_inches": round(section.page_height.inches, 3),
            "orientation": str(section.orientation),
            "top_margin_inches": round(section.top_margin.inches, 3),
            "bottom_margin_inches": round(section.bottom_margin.inches, 3),
            "left_margin_inches": round(section.left_margin.inches, 3),
            "right_margin_inches": round(section.right_margin.inches, 3),
        }
    )

with zipfile.ZipFile(docx_path) as archive:
    names = set(archive.namelist())
    document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    comments_xml = (
        archive.read("word/comments.xml").decode("utf-8", errors="ignore")
        if "word/comments.xml" in names
        else ""
    )
    footnotes_xml = (
        archive.read("word/footnotes.xml").decode("utf-8", errors="ignore")
        if "word/footnotes.xml" in names
        else ""
    )
    media_files = sorted(name for name in names if name.startswith("word/media/"))

report = {
    "page_count": len(pages),
    "paragraph_count": len(doc.paragraphs),
    "nonempty_paragraph_count": len(paragraphs),
    "table_count": len(tables),
    "inline_shape_count": len(doc.inline_shapes),
    "media_file_count": len(media_files),
    "section_count": len(doc.sections),
    "sections": sections,
    "style_counts": dict(style_counts.most_common()),
    "tracked_insertions": len(re.findall(r"<w:ins(?:\s|>)", document_xml)),
    "tracked_deletions": len(re.findall(r"<w:del(?:\s|>)", document_xml)),
    "comment_count": len(re.findall(r"<w:comment\b", comments_xml)),
    "footnote_count": max(0, len(re.findall(r"<w:footnote\b", footnotes_xml)) - 2),
    "media_files": media_files,
    "paragraphs": paragraphs,
    "tables": tables,
}
(out_dir / "structure.json").write_text(
    json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
)

# Render every PDF page and combine eight readable thumbnails per contact sheet.
pdf = pdfium.PdfDocument(str(pdf_path))
thumb_w, thumb_h = 330, 430
cols, rows_per_sheet = 4, 2
gap, label_h = 12, 26
per_sheet = cols * rows_per_sheet
for sheet_start in range(0, len(pdf), per_sheet):
    sheet = Image.new(
        "RGB",
        (
            cols * thumb_w + (cols + 1) * gap,
            rows_per_sheet * (thumb_h + label_h) + (rows_per_sheet + 1) * gap,
        ),
        "#cfd6da",
    )
    draw = ImageDraw.Draw(sheet)
    for offset in range(per_sheet):
        page_index = sheet_start + offset
        if page_index >= len(pdf):
            break
        page = pdf[page_index]
        bitmap = page.render(scale=0.65)
        image = bitmap.to_pil().convert("RGB")
        image.thumbnail((thumb_w, thumb_h))
        col = offset % cols
        row = offset // cols
        x = gap + col * thumb_w + max(0, (thumb_w - image.width) // 2)
        y = gap + row * (thumb_h + label_h)
        sheet.paste(image, (x, y))
        draw.text((gap + col * thumb_w + 3, y + thumb_h + 4), f"Page {page_index + 1}", fill="#111111")
    first_page = sheet_start + 1
    last_page = min(sheet_start + per_sheet, len(pdf))
    sheet.save(out_dir / f"contact-{first_page:03d}-{last_page:03d}.png")

print(json.dumps({key: report[key] for key in report if key not in {"paragraphs", "tables", "media_files"}}, indent=2))
