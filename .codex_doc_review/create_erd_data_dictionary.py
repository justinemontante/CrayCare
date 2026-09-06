from __future__ import annotations

import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(r"C:\Users\SLUMDUNK\Desktop\Crayfish Flutter UI\craycare")
DBML = ROOT / "docs" / "craycare_erd.dbml"
PUP_TEMPLATE = ROOT / "docs" / "Chapters 1–3.docx"
OUTPUT = ROOT / "docs" / "Data Dictionary.docx"

spec = importlib.util.spec_from_file_location(
    "doc_helpers", ROOT / ".codex_doc_review" / "create_simple_data_dictionary.py"
)
helpers = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(helpers)


SPECIAL_DESCRIPTIONS = {
    "uid": "Unique identifier of the user account.",
    "owner_uid": "Unique identifier of the tank owner.",
    "tank_id": "Unique identifier of the related tank.",
    "batch_id": "Unique identifier of the related production batch.",
    "log_id": "Unique identifier of the related feeding log.",
    "full_name": "Complete name of the user.",
    "email": "Email address used for sign-in and account recovery.",
    "role": "Access role assigned to the user.",
    "status": "Current status of the record or operation.",
    "photo_url": "Link to the user's profile image.",
    "title": "Short title shown to the user.",
    "body": "Complete message shown to the user.",
    "action": "Description of the action that occurred.",
    "type": "Category or type of the recorded event.",
    "created_at": "Date and time when the record was created.",
    "updated_at": "Date and time when the record was last updated.",
    "recorded_at": "Date and time when the sensor data was recorded.",
    "logged_at": "Date and time when the event was recorded in the log.",
    "assigned_at": "Date and time when the hardware was assigned.",
    "issued_at": "Date and time when the command was created.",
    "expires_at": "Date and time when the command expires.",
    "sampling_date": "Date and time when sampling was performed.",
    "mortality_date": "Date when the mortality was discovered or recorded.",
    "harvest_date": "Date when harvesting was performed.",
    "stocking_date": "Date when the production batch was stocked.",
    "temperature": "Current water temperature reading.",
    "ph_level": "Current pH reading of the water.",
    "dissolved_oxygen": "Current dissolved oxygen reading.",
    "turbidity": "Current water turbidity reading.",
    "water_level": "Current pond water-level reading.",
    "feed_level": "Estimated percentage of feed remaining in the hopper.",
    "estimated_feed_grams": "Estimated amount of feed remaining in the hopper.",
    "min_value": "Lowest value of the configured normal range.",
    "max_value": "Highest value of the configured normal range.",
    "critical_value": "Value below which the feed level is considered critical.",
    "current_state": "Actual state reported by the ESP32.",
    "control_mode": "Operating mode selected by the owner.",
    "is_read": "Shows whether the user has opened or marked the notification as read.",
}


def split_options(text: str) -> list[str]:
    parts, current, quoted = [], [], False
    for char in text:
        if char == "'":
            quoted = not quoted
        if char == "," and not quoted:
            parts.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    if current:
        parts.append("".join(current).strip())
    return parts


def parse_options(text: str | None):
    if not text:
        return "Optional", ""
    constraints = []
    note = ""
    for option in split_options(text):
        lower = option.lower()
        if lower == "pk":
            constraints.append("Primary Key")
        elif lower == "unique":
            constraints.append("Unique")
        elif lower.startswith("ref:"):
            target = option.split(":", 1)[1].strip().lstrip(">-< ")
            constraints.append(f"Foreign Key: {target}")
        elif lower.startswith("default:"):
            constraints.append("Default: " + option.split(":", 1)[1].strip())
        elif lower.startswith("note:"):
            note = option.split(":", 1)[1].strip().strip("'")
        else:
            constraints.append(option)
    return "; ".join(constraints) if constraints else "Optional", note


def format_example_for(table: str, field: str, dtype: str, note: str) -> str:
    table_exact = {
        ("users", "status"): "active",
        ("feeder_status", "status"): "completed",
        ("feeder_logs", "status"): "completed",
        ("feeder_logs", "type"): "auto",
        ("actuator_logs", "type"): "auto",
        ("hardware_assignments", "id"): "currentOwner",
        ("sensor_thresholds", "id"): "temperature",
        ("sensor_ingestion", "id"): "current",
        ("machine_learning_assessments", "id"): "20260831T143000",
        ("sampling_records", "id"): "baseline",
        ("feeder_schedules", "id"): "schedule_001",
        ("feeder_logs", "id"): "feedlog_001",
        ("feeder_commands", "id"): "command_001",
        ("notifications", "id"): "notification_001",
    }
    if (table, field) in table_exact:
        return table_exact[(table, field)]
    exact = {
        "full_name": "Juan Dela Cruz",
        "current_batch_id": "batch_2026_001",
        "role": "owner",
        "status": "active",
        "batch_status": "active",
        "sensor_name": "temperature",
        "device_id": "pump",
        "control_mode": "auto",
        "current_state": "on",
        "command_type": "feed_now",
        "actuator_type": "pump",
        "action": "Automatic feeding completed",
        "status_reason": "Insufficient feed",
        "last_outcome": "completed",
        "schedule_time": "6:00 AM",
        "verification_note": "Feed level decreased after dispensing",
        "routing_reason": "Assignment mismatch",
        "source": "RandomForestClassifier",
        "model_algorithm": "RandomForestClassifier",
        "model_version": "wqa_rf_v1",
        "training_data_origin": "synthetic_prototype",
        "label_origin": "deterministic_threshold_assessment",
        "driver_label": "Dissolved Oxygen",
        "driver_unit": "mg/L",
        "problem": "Dissolved oxygen is decreasing",
        "insight": "Dissolved oxygen decreased during the latest readings.",
        "notes": "Operating normally",
        "title": "Low Feed Level",
        "body": "Feed level is low at 18%. Refill soon.",
        "notif_type": "warning",
        "level": "Moderate",
        "model_level": "Moderate",
        "data_status": "ready",
        "driver": "DO",
        "time": "06:00",
        "ampm": "AM",
        "time_value": "360",
        "days": "1111111",
        "grams": "20.00",
        "feed_level": "68.00",
        "ph_level": "7.30",
        "dissolved_oxygen": "5.60",
        "temperature": "27.20",
        "turbidity": "31.00",
        "water_level": "18.00",
        "is_done": "false",
        "enabled": "true",
        "is_read": "false",
        "is_baseline": "false",
        "is_initialized": "true",
        "safety_override": "false",
        "level_change_detected": "true",
        "effective_at_ms": "1788167400000",
        "last_occurrence_at": "1788167400000",
        "occurrence_at": "1788167400000",
        "source_assignment_at_ms": "1788167400000",
        "captured_at_ms": "1788167400000",
        "ts_epoch": "1788167400",
        "amount_basis": "servo_cycle_estimate",
        "routing_status": "quarantined",
        "id": "record_001",
        "uid": "aB3dE6fG9hJ2kL5mN8pQ1rS4tU7v",
        "owner_uid": "aB3dE6fG9hJ2kL5mN8pQ1rS4tU7v",
        "issued_by": "aB3dE6fG9hJ2kL5mN8pQ1rS4tU7v",
        "assigned_by": "admin_uid_001",
        "tank_id": "tank_001",
        "batch_id": "batch_2026_001",
        "log_id": "feedlog_001",
        "command_id": "command_001",
        "schedule_key": "schedule_001",
        "email": "juan@example.com",
        "photo_url": "https://example.com/profile.jpg",
        "hardware_id": "ESP32_CRAYCARE_001",
        "source_tank_id": "tank_001",
        "source_owner_uid": "aB3dE6fG9hJ2kL5mN8pQ1rS4tU7v",
        "critical_value": "10.00",
        "hopper_capacity_grams": "1000.00",
        "min_value": "24.00",
        "max_value": "30.00",
        "estimated_feed_grams": "680.00",
        "requested_grams": "20.00",
        "estimated_dispensed_grams": "20.00",
        "estimated_available_grams": "680.00",
        "feed_level_before": "68.00",
        "feed_level_after": "66.00",
        "confidence": "86",
        "avg_body_weight": "25.50",
        "avg_body_length": "9.20",
        "total_weight": "255.00",
        "total_length": "92.00",
        "biomass": "2550.00",
        "harvest_weight_grams": "5000.00",
        "total_weight_kg": "5.250",
        "abw_grams": "35.50",
    }
    if field in exact:
        return exact[field]
    if field.startswith("temp_"):
        return "27.20"
    if field.startswith("pH_"):
        return "7.30"
    if field.startswith("DO_"):
        return "5.60"
    if field.startswith("turbidity_"):
        return "31.00"
    if field.startswith("waterLevel_"):
        return "18.00"
    if field.endswith("_at") or field.endswith("_date") or dtype == "timestamp":
        return "2026-08-31 14:30:00"
    if dtype == "boolean":
        return "true"
    if dtype == "int":
        return "20"
    if dtype == "bigint":
        return "1788167400000"
    if dtype.startswith("decimal"):
        return "25.50"
    if dtype in {"varchar", "text"}:
        return "Sample text"
    return dtype


def description_for(table: str, field: str, note: str) -> str:
    if note:
        return helpers.simple(note.rstrip(".")) + "."
    if field in SPECIAL_DESCRIPTIONS:
        return SPECIAL_DESCRIPTIONS[field]
    match = re.fullmatch(r"(.+)_(min|max|avg|sum|count)", field)
    if match:
        sensor = match.group(1).replace("waterLevel", "water level").replace("DO", "dissolved oxygen").replace("pH", "pH").replace("_", " ")
        kind = {
            "min": "Lowest",
            "max": "Highest",
            "avg": "Average",
            "sum": "Total used to calculate the average for",
            "count": "Number of valid readings included for",
        }[match.group(2)]
        if match.group(2) in {"sum", "count"}:
            return f"{kind} {sensor}."
        return f"{kind} recorded {sensor} value."
    label = field.replace("_", " ")
    if field.startswith("is_"):
        return f"Shows whether {label[3:]} is true or false."
    return f"Stores the {label}."


def parse_dbml():
    tables = []
    current = None
    depth = 0
    for raw in DBML.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        table_match = re.match(r"Table\s+([A-Za-z_][A-Za-z0-9_]*)", line)
        if table_match and line.endswith("{"):
            current = {"name": table_match.group(1), "columns": []}
            tables.append(current)
            depth = 1
            continue
        if current is None:
            continue
        if line.startswith("Indexes") and line.endswith("{"):
            depth += 1
            continue
        if line == "}":
            depth -= 1
            if depth == 0:
                current = None
            continue
        if depth != 1 or not line or line.startswith("//") or line.startswith("Note:"):
            continue
        column_match = re.match(r"([A-Za-z_][A-Za-z0-9_]*)\s+([^\s\[]+)(?:\s+\[(.*)\])?$", line)
        if not column_match:
            continue
        field, dtype, options = column_match.groups()
        constraints, note = parse_options(options)
        description = description_for(current["name"], field, note)
        if "Primary Key" in constraints:
            description += " Primary key of the table."
        foreign = re.search(r"Foreign Key: ([^;]+)", constraints)
        if foreign:
            description += f" References {foreign.group(1)}."
        current["columns"].append(
            (field, dtype, description, format_example_for(current["name"], field, dtype, note))
        )
    return tables


def build():
    tables = parse_dbml()
    doc = Document(PUP_TEMPLATE)
    body = doc._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("Data Dictionary")
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.bold = True

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.add_run(
        "This data dictionary presents the relational database structure shown in the CrayCare Entity-Relationship Diagram (ERD). "
        "Each table lists its exact table name and column names, data type, description, and the format or example value used by the current database."
    )
    intro2 = doc.add_paragraph()
    intro2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro2.add_run(
        "The table and column names are written exactly as defined in the ERD. Primary keys identify each record, while foreign keys connect related records such as users, tanks, batches, sensor data, feeding activities, and notifications."
    )

    for item in tables:
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(10)
        heading.paragraph_format.space_after = Pt(6)
        heading_run = heading.add_run(f"Table: {item['name']}")
        heading_run.font.name = "Arial"
        heading_run.font.size = Pt(10.5)
        heading_run.bold = True

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = helpers.WD_TABLE_ALIGNMENT.CENTER
        headers = ["Field Name", "Data Type", "Description", "Format / Example"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            helpers.style_cell(cell, header=True)
        helpers.cant_split(table.rows[0])
        helpers.set_repeat_header(table.rows[0])

        for field, dtype, description, constraints in item["columns"]:
            cells = table.add_row().cells
            for cell, value in zip(cells, (field, dtype, description, constraints)):
                cell.text = value
                helpers.style_cell(cell)
            helpers.cant_split(table.rows[-1])
        helpers.set_table_widths(
            table, [Inches(1.10), Inches(0.85), Inches(2.60), Inches(1.35)]
        )
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    doc.core_properties.title = "Data Dictionary"
    doc.core_properties.subject = "CrayCare ERD Data Dictionary"
    doc.core_properties.author = "CrayCare"
    try:
        doc.save(OUTPUT)
        saved_to = OUTPUT
    except PermissionError:
        saved_to = OUTPUT.with_name("Data Dictionary.pending.docx")
        doc.save(saved_to)
    print(saved_to)
    print(f"tables={len(tables)}")
    print(f"fields={sum(len(table['columns']) for table in tables)}")


if __name__ == "__main__":
    build()
